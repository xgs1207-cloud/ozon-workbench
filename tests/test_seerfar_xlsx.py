"""Seerfar 导出表导入测试：真实列名映射、值解析、入库与打分（含真实文件冒烟）。

真实表头来自用户 2026-10-05 的导出（24 列）；测试用**合成数据**构造同结构表格，
只在文件存在时对真实表做一次"结构冒烟"（不复制任何业务数据进仓库）。
"""

from __future__ import annotations

import json
import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from collector.seerfar_xlsx import (  # noqa: E402
    COLUMN_ALIASES,
    SeerfarError,
    clean_keyword,
    import_xlsx,
    parse_measure,
    parse_percent,
    read_seerfar_xlsx,
    split_category,
)
from keyword_library import store as keyword_store  # noqa: E402

REAL_EXPORT = pathlib.Path(r"D:\AI作图\Seerfar-Market20261005_2000.xlsx")

#: 用户导出表的真实表头顺序
HEADER = [
    "排名", "关键词", "类目", "关键词相关产品", "销售方式", "平均价格", "销量", "销售额",
    "月搜热度", "月搜增长", "商品数", "竞对数", "竞品数", "加购人数", "加购率", "转化率",
    "评论数", "评分", "市场空间", "商品可见度", "转化集中度", "重量", "体积", "退货取消率",
]


def row(
    *,
    rank: int,
    keyword: str,
    category: str,
    price: str = "1643₽",
    heat: float = 110906,
    growth: str = "6.54%",
    products: float = 10827,
    competitors: float = 206,
    rivals: float = 1800,
    weight: str = "782 g",
    volume: str = "4.97 L",
    return_rate: str = "12.31%",
) -> list:
    return [
        rank, keyword, category, "https://ir.ozone.ru/s3/multimedia-1-f/wc300/1.jpg", "跨境卖家可售",
        price, 1473.0, "1559210₽", heat, growth, products, competitors, rivals, 36094.0, "32.60%",
        "11.85%", 119295.0, 4.9, 62.0, 49.5, "34.34%", weight, volume, return_rate,
    ]


def write_xlsx(path: pathlib.Path, rows: list[list]) -> pathlib.Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(HEADER)
    for item in rows:
        sheet.append(item)
    workbook.save(path)
    return path


class ValueParsingTests(unittest.TestCase):
    def test_clean_keyword_strips_chinese_gloss_and_newlines(self):
        self.assertEqual(clean_keyword("простынь на резинке 160х200\n(橡胶枕160x200)"), "простынь на резинке 160х200")
        self.assertEqual(clean_keyword("простынь\n(被单)"), "простынь")
        self.assertEqual(clean_keyword("термос (保温杯) 500 мл"), "термос 500 мл")
        self.assertEqual(clean_keyword("  простынь  "), "простынь")

    def test_parse_percent(self):
        self.assertEqual(parse_percent("6.54%"), 6.54)
        self.assertEqual(parse_percent("12.31%"), 12.31)
        self.assertEqual(parse_percent(4.9), 4.9)
        self.assertIsNone(parse_percent(None))

    def test_parse_measure(self):
        self.assertEqual(parse_measure("782 g"), (782.0, "g"))
        self.assertEqual(parse_measure("4.97 L"), (4.97, "L"))
        self.assertEqual(parse_measure(1.5), (1.5, None))

    def test_split_category(self):
        self.assertEqual(split_category("床单\nПростыня"), ("床单", "Простыня"))
        self.assertEqual(split_category("床单"), ("床单", None))


class ColumnMappingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = pathlib.Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_all_24_real_columns_are_mapped(self):
        path = write_xlsx(
            self.root / "s.xlsx",
            [row(rank=1, keyword="термос 500 мл\n(保温杯)", category="保温杯\nТермос")],
        )
        parsed = read_seerfar_xlsx(path)
        self.assertEqual(parsed["unmapped_columns"], [])
        self.assertEqual(len(parsed["columns"]), len(COLUMN_ALIASES))
        record = parsed["records"][0]
        self.assertEqual(record["keyword"], "термос 500 мл")
        self.assertEqual(record["search_volume"], 110906)
        self.assertEqual(record["competitor_count"], 206)  # 默认口径：竞对数
        self.assertEqual(record["trend"], 6.54)
        extra = record["extra"]
        self.assertEqual(extra["category_name_zh"], "保温杯")
        self.assertEqual(extra["category_name_ru"], "Термос")
        self.assertEqual(extra["product_count"], 10827)
        self.assertEqual(extra["rival_count"], 1800)
        self.assertEqual(extra["weight_g"], 782)
        self.assertEqual(extra["volume_l"], 4.97)
        self.assertEqual(extra["return_rate_percent"], 12.31)
        self.assertEqual(extra["revenue_rub"], 1559210)
        self.assertEqual(extra["competition_field"], "竞对数")

    def test_competition_field_can_be_switched(self):
        path = write_xlsx(self.root / "s.xlsx", [row(rank=1, keyword="термос", category="保温杯\nТермос")])
        parsed = read_seerfar_xlsx(path, competition_field="商品数")
        self.assertEqual(parsed["records"][0]["competitor_count"], 10827)
        self.assertEqual(parsed["records"][0]["extra"]["competition_field"], "商品数")
        with self.assertRaises(SeerfarError):
            read_seerfar_xlsx(path, competition_field="不存在的列")

    def test_missing_required_column_reports_headers(self):
        import openpyxl

        path = self.root / "broken.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["排名", "类目"])  # 缺关键词与月搜热度
        sheet.append([1, "床单"])
        workbook.save(path)
        with self.assertRaises(SeerfarError) as ctx:
            read_seerfar_xlsx(path)
        message = str(ctx.exception)
        self.assertIn("关键词", message)
        self.assertIn("月搜热度", message)

    def test_unknown_columns_are_reported_and_kept_nothing_lost(self):
        import openpyxl

        path = self.root / "extra.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["排名", "关键词", "月搜热度", "竞对数", "未来的新指标"])
        sheet.append([1, "термос", 1000, 10, "42"])
        workbook.save(path)
        parsed = read_seerfar_xlsx(path)
        self.assertEqual(parsed["unmapped_columns"], ["未来的新指标"])
        self.assertEqual(len(parsed["records"]), 1)

    def test_blank_rows_are_skipped(self):
        path = write_xlsx(
            self.root / "s.xlsx",
            [row(rank=1, keyword="термос", category="保温杯\nТермос"), [None] * len(HEADER)],
        )
        self.assertEqual(len(read_seerfar_xlsx(path)["records"]), 1)

    def test_empty_keyword_row_is_reported_not_silently_dropped(self):
        path = write_xlsx(
            self.root / "s.xlsx",
            [row(rank=1, keyword="термос", category="保温杯\nТермос"), row(rank=2, keyword="   ", category="保温杯\nТермос")],
        )
        parsed = read_seerfar_xlsx(path)
        self.assertEqual(len(parsed["records"]), 1)
        self.assertEqual(len(parsed["skipped_rows"]), 1)


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = pathlib.Path(self.tmp.name)
        self.library = self.root / "keyword-library"
        self.xlsx = write_xlsx(
            self.root / "seerfar.xlsx",
            [
                row(rank=1, keyword="простынь на резинке 160х200\n(橡胶床单)", category="床单\nПростыня", heat=110906, competitors=206),
                row(rank=2, keyword="простынь\n(被单)", category="床单\nПростыня", heat=88185, competitors=249),
                row(rank=3, keyword="термос 500 мл\n(保温杯)", category="保温杯\nТермос", heat=51000, competitors=90),
            ],
        )

    def tearDown(self):
        self.tmp.cleanup()

    def test_import_groups_by_category_and_scores(self):
        summary = import_xlsx(self.xlsx, self.library, competition_field="竞对数")
        self.assertTrue(summary["ok"])
        self.assertEqual(summary["rows_parsed"], 3)
        self.assertEqual(summary["keywords_imported"], 3)
        self.assertEqual(len(summary["categories"]), 2)
        self.assertTrue(any("合成主键" in item for item in summary["warnings"]))

        # 两个类目各自成文件；打分只在同组内比较
        files = sorted(path.name for path in self.library.glob("*.jsonl"))
        self.assertEqual(len(files), 2)
        records = keyword_store.load_all(self.library)
        self.assertEqual(len(records), 3)
        for record in records:
            self.assertIsNotNone(record["score"])
            self.assertEqual(record["source"], "seerfar")
            self.assertEqual(record["search_volume"], record["extra"]["search_heat"])
            self.assertTrue(record["extra"]["category_name_ru"])

    def test_reimport_is_idempotent_and_updates(self):
        import_xlsx(self.xlsx, self.library)
        second = import_xlsx(self.xlsx, self.library)
        self.assertEqual(second["created"], 0)
        self.assertEqual(second["updated"], 3)
        self.assertEqual(len(keyword_store.load_all(self.library)), 3)

    def test_explicit_category_ids_are_used(self):
        summary = import_xlsx(self.xlsx, self.library, category_id="1001", type_id="2001")
        self.assertEqual(summary["categories"], {"1001:2001": 3})
        self.assertFalse(any("合成主键" in item for item in summary["warnings"]))
        self.assertTrue((self.library / "1001-2001.jsonl").is_file())

    def test_status_promotion_after_import(self):
        import_xlsx(self.xlsx, self.library, category_id="1001", type_id="2001")
        qualified = keyword_store.query(
            self.library, category_id="1001", type_id="2001", status=keyword_store.STATUS_QUALIFIED
        )
        candidates = keyword_store.query(
            self.library, category_id="1001", type_id="2001", status=keyword_store.STATUS_CANDIDATE
        )
        self.assertEqual(len(qualified) + len(candidates), 3)
        # 高分低竞争的那条应当入选（或至少拿到最高分）
        ranked = keyword_store.query(self.library, category_id="1001", type_id="2001", order="score")
        self.assertIsNotNone(ranked[0]["score"])
        self.assertGreaterEqual(ranked[0]["score"], ranked[-1]["score"])


@unittest.skipUnless(REAL_EXPORT.is_file(), "用户的真实导出表不在本机")
class RealExportSmokeTests(unittest.TestCase):
    """只验证结构能读懂（不把业务数据写进任何测试产物）。"""

    def test_real_export_structure_is_fully_understood(self):
        parsed = read_seerfar_xlsx(REAL_EXPORT)
        self.assertEqual(parsed["sheets"], ["Data"])
        self.assertGreater(len(parsed["records"]), 500)
        self.assertEqual(parsed["unmapped_columns"], [])
        self.assertEqual(parsed["skipped_rows"], [])
        sample = parsed["records"][0]
        self.assertTrue(sample["keyword"])
        self.assertIsNotNone(sample["search_volume"])
        self.assertIsNotNone(sample["competitor_count"])
        self.assertTrue(sample["extra"]["category_name_ru"])

    def test_real_export_imports_into_temp_library(self):
        with tempfile.TemporaryDirectory() as directory:
            summary = import_xlsx(REAL_EXPORT, pathlib.Path(directory) / "lib")
        self.assertGreater(summary["keywords_imported"], 500)
        self.assertGreaterEqual(len(summary["categories"]), 1)
        self.assertTrue(summary["rescored"])


if __name__ == "__main__":
    unittest.main()
